#!/usr/bin/env python3
"""Offline functional verification for the managed Python environment."""

from __future__ import annotations

import argparse
import importlib
import importlib.metadata
import json
import os
import platform
import re
import subprocess
import sys
import tempfile
import traceback
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.build.requirements_lock import LockError, locked_versions  # noqa: E402

IMPORT_NAMES = {
    "beautifulsoup4": "bs4",
    "faker": "faker",
    "flask": "flask",
    "matplotlib": "matplotlib",
    "mysql-connector-python": "mysql.connector",
    "networkx": "networkx",
    "numpy": "numpy",
    "openpyxl": "openpyxl",
    "paho-mqtt": "paho.mqtt.client",
    "pandas": "pandas",
    "scikit-learn": "sklearn",
    "seaborn": "seaborn",
    "selenium": "selenium",
    "webdriver-manager": "webdriver_manager",
}


class VerificationError(RuntimeError):
    """Raised when a required environment contract is not satisfied."""


def _canonical_path(value: str | os.PathLike[str]) -> str:
    return os.path.normcase(os.path.realpath(os.fspath(value)))


def verify_runtime(expected_python: str, expected_executable: str | None) -> None:
    if platform.python_version() != expected_python:
        raise VerificationError(
            "Python version mismatch: "
            f"expected={expected_python} actual={platform.python_version()}"
        )
    if platform.architecture()[0] != "64bit":
        raise VerificationError(f"Python architecture is not 64bit: {platform.architecture()[0]}")
    if sys.implementation.name != "cpython":
        raise VerificationError(f"Python implementation is not CPython: {sys.implementation.name}")
    if sys.prefix == sys.base_prefix:
        raise VerificationError("Interpreter is not running inside a virtual environment")
    if expected_executable and _canonical_path(sys.executable) != _canonical_path(
        expected_executable
    ):
        raise VerificationError(
            f"Python executable mismatch: expected={expected_executable} actual={sys.executable}"
        )


def verify_packages(requirements: Path) -> dict[str, str]:
    expected = locked_versions(requirements)
    installed: dict[str, str] = {}
    for name, expected_version in sorted(expected.items()):
        try:
            actual_version = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError as exc:
            raise VerificationError(f"Locked package is not installed: {name}") from exc
        if actual_version != expected_version:
            raise VerificationError(
                f"Package version mismatch: {name} "
                f"expected={expected_version} actual={actual_version}"
            )
        installed[name] = actual_version
    for package, module in IMPORT_NAMES.items():
        if package not in expected:
            raise VerificationError(f"Required direct package missing from lock: {package}")
        try:
            importlib.import_module(module)
        except Exception as exc:
            raise VerificationError(f"Import failed: {module}: {exc}") from exc
    completed = subprocess.run(
        [sys.executable, "-I", "-m", "pip", "check"],
        check=False,
        capture_output=True,
        text=True,
        timeout=120,
    )
    if completed.returncode:
        raise VerificationError(
            f"pip check failed ({completed.returncode}): {completed.stdout} {completed.stderr}"
        )
    return installed


def verify_bootstrap(requirements: Path) -> dict[str, str]:
    expected = locked_versions(requirements)
    installed: dict[str, str] = {}
    for name, expected_version in sorted(expected.items()):
        try:
            actual_version = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError as exc:
            raise VerificationError(f"Bootstrap package is not installed: {name}") from exc
        if actual_version != expected_version:
            raise VerificationError(
                f"Bootstrap version mismatch: {name} "
                f"expected={expected_version} actual={actual_version}"
            )
        installed[name] = actual_version
    return installed


def verify_immutable_distribution_set(
    packages: dict[str, str], bootstrap_tooling: dict[str, str]
) -> None:
    expected = set(packages) | set(bootstrap_tooling)
    actual = {
        re.sub(r"[-_.]+", "-", name).lower()
        for distribution in importlib.metadata.distributions()
        if (name := distribution.metadata.get("Name"))
    }
    extras = sorted(actual - expected)
    missing = sorted(expected - actual)
    if extras or missing:
        raise VerificationError(
            f"Installed distribution set drifted: extras={extras} missing={missing}"
        )


_WINDOWS_INVALID_FILENAME_CHARACTERS = frozenset('<>:"/\\|?*')
_WINDOWS_RESERVED_FILENAMES = frozenset(
    {
        "AUX",
        "CON",
        "NUL",
        "PRN",
        *(f"COM{number}" for number in range(1, 10)),
        *(f"LPT{number}" for number in range(1, 10)),
    }
)


def _is_safe_entrypoint_name(name: str) -> bool:
    if not name or name in {".", ".."} or name[-1] in {" ", "."}:
        return False
    if any(
        ord(character) < 32 or character in _WINDOWS_INVALID_FILENAME_CHARACTERS
        for character in name
    ):
        return False
    return name.split(".", maxsplit=1)[0].upper() not in _WINDOWS_RESERVED_FILENAMES


def _entrypoint_launchers() -> list[tuple[str, Path]]:
    console_target = Path(sys.executable)
    gui_target = console_target.with_name("pythonw.exe")
    launchers: dict[str, Path] = {}
    for distribution in importlib.metadata.distributions():
        distribution_name = distribution.metadata.get("Name", "")
        canonical_distribution = re.sub(r"[-_.]+", "-", distribution_name).lower()
        for entry_point in distribution.entry_points:
            if entry_point.group not in {"console_scripts", "gui_scripts"}:
                continue
            name = entry_point.name
            if not _is_safe_entrypoint_name(name):
                raise VerificationError(f"Unsafe command entry-point name: {name!r}")
            target = console_target if entry_point.group == "console_scripts" else gui_target
            existing_target = launchers.get(name)
            if existing_target is not None and existing_target != target:
                raise VerificationError(f"Ambiguous command entry-point launcher: {name}")
            launchers[name] = target
            if canonical_distribution == "pip" and name == "pip":
                launchers[f"pip{sys.version_info.major}"] = console_target
                launchers[f"pip{sys.version_info.major}.{sys.version_info.minor}"] = console_target
    if not launchers:
        raise VerificationError("No command entry-point launchers were found")
    return sorted(launchers.items())


def verify_entrypoint_launchers() -> None:
    scripts_directory = Path(sys.executable).parent
    for name, target in _entrypoint_launchers():
        target_bytes = os.fsencode(target)
        shebangs = {
            b"#!" + target_bytes + b"\n",
            b'#!"' + target_bytes + b'"\n',
        }
        launcher = scripts_directory / f"{name}.exe"
        try:
            content = launcher.read_bytes()
        except OSError as exc:
            raise VerificationError(f"Command launcher is missing: {launcher}") from exc
        if not any(shebang in content for shebang in shebangs):
            raise VerificationError(f"Command launcher targets the wrong interpreter: {launcher}")


def smoke_scientific_and_files() -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    import pandas as pd
    from openpyxl import Workbook, load_workbook
    from sklearn.linear_model import LinearRegression

    values = np.array([1.0, 2.0, 3.0])
    if float(values.mean()) != 2.0:
        raise VerificationError("NumPy mean smoke test failed")
    frame = pd.DataFrame({"value": values})
    if float(frame["value"].sum()) != 6.0:
        raise VerificationError("Pandas sum smoke test failed")

    with tempfile.TemporaryDirectory(prefix="python-runtime-verification-") as temporary:
        temporary_path = Path(temporary)
        workbook_path = temporary_path / "verification.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "verified"
        workbook.save(workbook_path)
        if load_workbook(workbook_path, read_only=True).active["A1"].value != "verified":
            raise VerificationError("OpenPyXL round-trip smoke test failed")

        image_path = temporary_path / "verification.png"
        figure, axis = plt.subplots()
        axis.plot([0, 1], [0, 1])
        figure.savefig(image_path)
        plt.close(figure)
        if not image_path.is_file() or image_path.stat().st_size == 0:
            raise VerificationError("Matplotlib rendering smoke test failed")

    model = LinearRegression().fit([[0.0], [1.0], [2.0]], [0.0, 2.0, 4.0])
    if abs(float(model.predict([[3.0]])[0]) - 6.0) > 1e-8:
        raise VerificationError("scikit-learn training smoke test failed")


def smoke_services_and_utilities() -> None:
    import mysql.connector
    import networkx as nx
    import paho.mqtt.client as mqtt
    from bs4 import BeautifulSoup
    from faker import Faker
    from flask import Flask

    application = Flask("runtime-verification")

    @application.get("/health")
    def health() -> tuple[dict[str, bool], int]:
        return {"ok": True}, 200

    response = application.test_client().get("/health")
    if response.status_code != 200 or response.get_json() != {"ok": True}:
        raise VerificationError("Flask test-client smoke test failed")

    graph = nx.Graph()
    graph.add_edge("a", "b")
    if nx.shortest_path(graph, "a", "b") != ["a", "b"]:
        raise VerificationError("NetworkX smoke test failed")

    Faker.seed(42)
    if not Faker("en_US").name():
        raise VerificationError("Faker smoke test failed")
    if BeautifulSoup("<p>verified</p>", "html.parser").p.text != "verified":
        raise VerificationError("BeautifulSoup smoke test failed")

    connection = mysql.connector.MySQLConnection()
    if connection.is_connected():
        raise VerificationError("MySQL connector unexpectedly opened a connection")
    connection.close()

    callback_api = getattr(mqtt, "CallbackAPIVersion", None)
    client = mqtt.Client(callback_api.VERSION2) if callback_api else mqtt.Client()
    if client is None:
        raise VerificationError("Paho MQTT client construction failed")

    importlib.import_module("selenium")
    importlib.import_module("webdriver_manager")


def verify_manifest(
    manifest_path: Path,
    packages: dict[str, str],
    bootstrap_tooling: dict[str, str],
    expected_python: str,
) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("python_version") != expected_python:
        raise VerificationError("Installed manifest Python version does not match")
    manifest_executable = manifest.get("python_executable")
    if not isinstance(manifest_executable, str) or _canonical_path(
        manifest_executable
    ) != _canonical_path(sys.executable):
        raise VerificationError("Installed manifest executable does not match the environment")
    if manifest.get("verification_status") != "passed":
        raise VerificationError("Installed manifest is not marked as verified")
    manifest_packages = manifest.get("packages")
    if not isinstance(manifest_packages, dict) or manifest_packages != packages:
        raise VerificationError("Installed manifest package versions do not match the environment")
    if manifest.get("bootstrap_tooling") != bootstrap_tooling:
        raise VerificationError(
            "Installed manifest bootstrap versions do not match the environment"
        )


def run_verification(
    requirements: Path,
    bootstrap_requirements: Path,
    expected_python: str,
    expected_executable: str | None,
    manifest: Path | None,
    smoke_tests: tuple[Callable[[], None], ...] = (
        smoke_scientific_and_files,
        smoke_services_and_utilities,
    ),
) -> dict[str, object]:
    verify_runtime(expected_python, expected_executable)
    packages = verify_packages(requirements)
    bootstrap_tooling = verify_bootstrap(bootstrap_requirements)
    verify_immutable_distribution_set(packages, bootstrap_tooling)
    verify_entrypoint_launchers()
    for smoke_test in smoke_tests:
        smoke_test()
    if manifest:
        verify_manifest(manifest, packages, bootstrap_tooling, expected_python)
    return {
        "schema_version": 1,
        "status": "passed",
        "verified_at": datetime.now(UTC).isoformat(),
        "python_version": platform.python_version(),
        "python_executable": sys.executable,
        "architecture": platform.architecture()[0],
        "packages": packages,
        "bootstrap_tooling": bootstrap_tooling,
        "checks": [
            "runtime",
            "locked-package-versions",
            "immutable-distribution-set",
            "entrypoint-launchers",
            "imports",
            "pip-check",
            "scientific-and-files",
            "services-and-utilities",
            "selenium-import-only",
        ],
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--requirements", type=Path, required=True)
    parser.add_argument("--bootstrap-requirements", type=Path, required=True)
    parser.add_argument("--expected-python", required=True)
    parser.add_argument("--expected-executable")
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    try:
        result = run_verification(
            args.requirements,
            args.bootstrap_requirements,
            args.expected_python,
            args.expected_executable,
            args.manifest,
        )
    except (VerificationError, LockError, OSError, ValueError, subprocess.SubprocessError) as exc:
        print(f"environment verification failed: {exc}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)
        return 40
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
